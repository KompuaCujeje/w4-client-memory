"""Старая таблица (Excel): организации без дублей по написанию, сделки без двойного счёта.

Без ИИ — правила явные: одна организация + один продукт = одна сделка, она за тем, кто завёл первым.
"""

from __future__ import annotations

import datetime as dt
import io

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .db import CLOSED, MANAGERS, PRODUCTS, STAGES, Event, Fact, Org, Project
from .names import org_key

# колонка таблицы → этап словаря
STAGE_COLS = {
    "создан": "знакомство",
    "потребность": "потребность",
    "предложение": "оценка и предложение",
    "согласование": "согласование",
    "договор": "договор",
}


def _d(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and v.strip():
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return dt.datetime.strptime(v.strip()[:10], fmt).date()
            except ValueError:
                pass
    return None


def _norm_stage(v: str) -> str:
    v = (v or "").strip().lower()
    for st in STAGES + [CLOSED]:
        if v == st:
            return st
    if "закрыт" in v:
        return CLOSED
    if "оценк" in v or "предлож" in v:
        return "оценка и предложение"
    if "работ" in v:
        return "в работе"
    return v or "знакомство"


def _norm_product(v: str) -> str:
    v = (v or "").strip().lower()
    for p in PRODUCTS:
        if v == p.lower():
            return p
    return (v or "").strip()


def find_org_table(s: Session, name: str) -> Org | None:
    # ключ без кавычек, формы собственности, дефисов и пробелов:
    # «ООО «Ритейл-Групп»» = «Ритейл Групп», «Стоматология "Улыбка"» = «Стоматология «Улыбка»»
    key = org_key(name)
    for o in s.scalars(select(Org)):
        if any(org_key(n) == key for n in [o.name] + list(o.aliases or [])):
            return o
    return None


def import_table(s: Session, ev: Event, data: bytes) -> None:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            for i, h in enumerate(head):
                if h.startswith(n):
                    return i
        return None

    c_no, c_client, c_prod, c_mgr = col("№", "no", "#"), col("клиент", "организац"), col("продукт"), col("менеджер")
    c_stage, c_budget, c_comment = col("этап"), col("бюджет"), col("коммент")
    stage_cols = {st: col(h) for h, st in STAGE_COLS.items()}
    if ev.date is None:
        ev.date = dt.date(2026, 6, 30)

    for n, r in enumerate(rows[1:], start=1):
        if not r or c_client is None or not r[c_client]:
            continue
        no = r[c_no] if c_no is not None and r[c_no] else n
        raw_name = str(r[c_client]).strip()
        product = _norm_product(str(r[c_prod] or "")) if c_prod is not None else ""
        manager = str(r[c_mgr] or "").strip() if c_mgr is not None else ""
        manager = next((m for m in MANAGERS if m.lower() in manager.lower()), manager)
        stage = _norm_stage(str(r[c_stage] or "")) if c_stage is not None else "знакомство"
        dates = {st: _d(r[i]) for st, i in stage_cols.items() if i is not None}
        dates = {k: v for k, v in dates.items() if v}
        budget = r[c_budget] if c_budget is not None else None
        comment = str(r[c_comment] or "").strip() if c_comment is not None else ""
        code = f"{ev.code}#{no}"

        row_text = " | ".join(f"{head[i]}: {v:%d.%m.%Y}" if isinstance(v, dt.datetime) else f"{head[i]}: {v}" for i, v in enumerate(r) if v not in (None, ""))
        last = max(dates.values()) if dates else ev.date
        s.add(Event(code=code, parent=ev.code, date=last, channel="таблица", author=manager, title=f"строка {no}", text=row_text, status="done"))

        org = find_org_table(s, raw_name)
        if org is None:
            org = Org(name=raw_name, aliases=[], first_event=code)
            s.add(org)
            s.flush()
        elif raw_name != org.name and raw_name not in (org.aliases or []):
            org.aliases = list(org.aliases or []) + [raw_name]

        dup = s.scalar(
            select(Project).where(Project.org_id == org.id, Project.product == product, Project.duplicate_of.is_(None))
        )
        proj = Project(
            org_id=org.id,
            title=product or "сделка",
            product=product,
            manager=manager,
            stage=stage,
            stage_dates={k: v.isoformat() for k, v in dates.items()},
            stage_history=[{"stage": k, "date": v.isoformat(), "event": code} for k, v in sorted(dates.items(), key=lambda x: x[1])],
            created=dates.get("знакомство"),
            source=code,
            comment=comment,
            close_reason=comment if stage == CLOSED else "",
        )
        s.add(proj)
        s.flush()
        if dup is not None:
            # одна организация + один продукт = одна сделка; она за тем, кто завёл первым
            first, second = (dup, proj) if (dup.created or dt.date.max) <= (proj.created or dt.date.max) else (proj, dup)
            second.duplicate_of = first.id
            s.add(Fact(org_id=org.id, project_id=first.id, kind="other", text=f"Дубль сделки: строка {second.source} повторяет {first.source} (завёл первым: {first.manager})", event_code=code, date=ev.date, channel="таблица"))
        if budget not in (None, ""):
            s.add(Fact(org_id=org.id, project_id=proj.id, kind="budget", text=f"Бюджет {budget} млн ₽ (таблица)", value=f"{budget} млн ₽", event_code=code, date=last, channel="таблица"))
        if comment:
            s.add(Fact(org_id=org.id, project_id=proj.id, kind="stage" if stage == CLOSED else "other", text=comment, event_code=code, date=last, channel="таблица"))
        ev.org_ids = list(dict.fromkeys(list(ev.org_ids or []) + [org.id]))
    s.commit()
